# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

BLUE = "0000FF"      # hardcoded inputs
BLACK = "000000"     # formulas
GREEN = "008000"     # cross-sheet links
WHITE = "FFFFFF"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
YEL_FILL = PatternFill("solid", fgColor="FFF2CC")
FONT = "Arial"

THB = '#,##0;(#,##0);-'
NUM = '#,##0;(#,##0);-'
PCT = '0.0%'

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def s(ws, cell, value, color=BLACK, bold=False, size=10, fill=None, fmt=None,
      align=None, italic=False, wrap=False, bd=False):
    c = ws[cell]
    c.value = value
    c.font = Font(name=FONT, color=color, bold=bold, size=size, italic=italic)
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    else: c.alignment = Alignment(vertical="center", wrap_text=wrap)
    if bd: c.border = border
    return c

MONTHS = ["ก.ย.","ต.ค.","พ.ย.","ธ.ค.","ม.ค.","ก.พ.","มี.ค.","เม.ย.","พ.ค.","มิ.ย.","ก.ค.","ส.ค."]
# seasonal demand as % of Jan peak (shape from original doc)
SEASON = [0.103,0.276,0.483,0.724,1.000,0.828,0.586,0.483,0.552,0.638,0.741,0.879]
# month columns B..M  (col 2..13)
MCOLS = [get_column_letter(2+i) for i in range(12)]

wb = Workbook()

# ============================================================ ASSUMPTIONS
a = wb.active
a.title = "Assumptions"
a.sheet_view.showGridLines = False
s(a, "A1", "TCAS AI Practice App — Financial Model (3 Scenarios)", WHITE, True, 14, HDR_FILL)
for col in "ABCDEF": a[f"{col}1"].fill = HDR_FILL
s(a, "A2", "ทุกตัวเลขสีน้ำเงิน = ปรับได้ / สีดำ = สูตร / สีเขียว = ลิงก์ข้ามชีต   |   หน่วย: บาท (THB)", "808080", italic=True, size=9)

# scenario headers
s(a, "A4", "Assumption", WHITE, True, 11, SUB_FILL); a["A4"].font = Font(name=FONT, bold=True, size=11)
hdrs = [("B","Conservative","C0504D"),("C","Base (แนะนำ)","1F4E78"),("D","Optimistic","2E7D32"),("E","หน่วย / หมายเหตุ","000000")]
s(a, "A4", "Assumption", WHITE, True, 11, PatternFill("solid", fgColor="404040"))
for col,name,clr in hdrs:
    s(a, f"{col}4", name, WHITE, True, 11, PatternFill("solid", fgColor=clr), align="center")

rows = [
    ("Peak Total Users (เดือน ม.ค.)", 1200, 2000, 3000, "คน — ยอดผู้ใช้สูงสุด (พีคฤดูสอบ)", NUM),
    ("Conversion Rate (free→paid)", 0.12, 0.20, 0.35, "% ของ user ที่จ่ายเงิน", PCT),
    ("ARPU (บาท/ผู้จ่าย/เดือน)", 129, 129, 129, "99×0.8 + 249×0.2", THB),
    ("Transaction + Store Fee", 0.18, 0.08, 0.03, "% — payment gateway + app store cut", PCT),
    ("CAC (บาท/ผู้ใช้ใหม่)", 20, 10, 4, "ต้นทุนหา user ใหม่ 1 คน — organic ช่วง launch (no paid mktg)", THB),
    ("Monthly Operating Cost", 11000, 9000, 7500, "บาท/เดือน — Claude Code sub (flat) + infra/hosting + gen/verify demand-driven", THB),
    ("Pre-Launch Capital", 25000, 12000, 6000, "บาท — infra setup + misc (seed bank gen ≈ 0 ผ่าน Claude Code + verify มือ)", THB),
    ("Free Period (เดือนแรกที่ไม่เก็บเงิน)", 0, 0, 0, "เดือน — reverse trial (trial สั้น → เก็บเงินตั้งแต่เดือน 1)", NUM),
]
r = 5
ROWMAP = {}
keys = ["peak","conv","arpu","fee","cac","opcost","prelaunch","free"]
for (label,cons,base,opt,note,fmt),key in zip(rows,keys):
    s(a, f"A{r}", label, BLACK, False, 10, bd=True)
    s(a, f"B{r}", cons, BLUE, False, 10, fmt=fmt, align="center", bd=True)
    s(a, f"C{r}", base, BLUE, False, 10, fmt=fmt, align="center", bd=True, fill=YEL_FILL)
    s(a, f"D{r}", opt, BLUE, False, 10, fmt=fmt, align="center", bd=True)
    s(a, f"E{r}", note, "808080", False, 9, italic=True)
    ROWMAP[key] = r
    r += 1

# pre-launch breakdown note
s(a, "A14", "Pre-Launch breakdown (Base ~12k):", BLACK, True, 9)
s(a, "A15", "Seed bank ~4k โจทย์: gen ผ่าน Claude Code (flat subscription, marginal cost ≈ 0) + verify = พี่ตรวจมือเอง (0 API)  |  เหลือแค่ infra setup + domain + misc", "808080", italic=True, size=9)
s(a, "A16", "Claude Code subscription (~$100-200/mo) เป็น fixed cost ที่จ่ายอยู่แล้ว ไม่ scale ตาม bank → ย้ายไปอยู่ใน Monthly Operating Cost ไม่ใช่ Pre-Launch", "808080", italic=True, size=9)

# seasonal curve
s(a, "A19", "Seasonal Demand Curve (% ของพีค ม.ค.)", WHITE, True, 11, PatternFill("solid", fgColor="404040"))
for col in "BCDEFGHIJKLM": a[f"{col}19"].fill = PatternFill("solid", fgColor="404040")
for i,m in enumerate(MONTHS):
    s(a, f"{MCOLS[i]}20", m, WHITE, True, 9, PatternFill("solid", fgColor="595959"), align="center")
s(a, "A20", "เดือน", WHITE, True, 9, PatternFill("solid", fgColor="595959"))
s(a, "A21", "% ของพีค", BLACK, False, 9)
for i,frac in enumerate(SEASON):
    s(a, f"{MCOLS[i]}21", frac, BLUE, False, 9, fmt=PCT, align="center", bd=True)

a.column_dimensions["A"].width = 34
for col in "BCD": a.column_dimensions[col].width = 14
a.column_dimensions["E"].width = 42
for col in "FGHIJKLM": a.column_dimensions[col].width = 8

def aref(key, scen_col):
    return f"Assumptions!${scen_col}${ROWMAP[key]}"

# ============================================================ SCENARIO SHEET
def build_scenario(name, scen_col, accent):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    s(ws, "A1", f"Scenario: {name}", WHITE, True, 13, PatternFill("solid", fgColor=accent))
    for c in ["A"]+MCOLS: ws[f"{c}1"].fill = PatternFill("solid", fgColor=accent)
    # month header
    s(ws, "A3", "เดือน", WHITE, True, 10, PatternFill("solid", fgColor="404040"))
    for i,m in enumerate(MONTHS):
        s(ws, f"{MCOLS[i]}3", m, WHITE, True, 10, PatternFill("solid", fgColor="404040"), align="center")

    peak = aref("peak", scen_col); conv = aref("conv", scen_col)
    arpu = aref("arpu", scen_col); fee = aref("fee", scen_col)
    cac = aref("cac", scen_col); opc = aref("opcost", scen_col)
    pre = aref("prelaunch", scen_col); free = aref("free", scen_col)

    labels = {
        4:"Season % of peak", 5:"Total Users", 6:"Charging? (1=เก็บเงิน)",
        7:"Paying Users", 8:"Gross Revenue (บาท)", 9:"Net Revenue (หลังค่าธรรมเนียม)",
        10:"New Users (acquired)", 11:"Marketing/CAC Cost", 12:"Operating Cost",
        13:"Total Cost", 14:"Net Profit / เดือน", 15:"Cumulative Cash (สะสม)",
    }
    fills = {8:SUB_FILL, 9:SUB_FILL, 14:YEL_FILL, 15:YEL_FILL}
    for row,lab in labels.items():
        bold = row in (5,7,9,14,15)
        s(ws, f"A{row}", lab, BLACK, bold, 10, fill=fills.get(row), bd=True)

    for i in range(12):
        c = MCOLS[i]; mi = i+1
        scol = c  # season is in same month column on Assumptions row 21
        # 4 season% (green link)
        s(ws, f"{c}4", f"=Assumptions!{scol}$21", GREEN, False, 9, fmt=PCT, align="center", bd=True)
        # 5 total users
        s(ws, f"{c}5", f"=ROUND({peak}*{c}4,0)", BLACK, True, 10, fmt=NUM, align="center", bd=True)
        # 6 charging flag
        s(ws, f"{c}6", f"=IF(COLUMN()-1>{free},1,0)", BLACK, False, 9, fmt='0', align="center", bd=True)
        # 7 paying
        s(ws, f"{c}7", f"=ROUND({c}5*{conv}*{c}6,0)", BLACK, True, 10, fmt=NUM, align="center", bd=True)
        # 8 gross rev
        s(ws, f"{c}8", f"={c}7*{arpu}", BLACK, False, 10, fmt=THB, align="center", bd=True, fill=SUB_FILL)
        # 9 net rev
        s(ws, f"{c}9", f"={c}8*(1-{fee})", BLACK, True, 10, fmt=THB, align="center", bd=True, fill=SUB_FILL)
        # 10 new users
        if i==0:
            s(ws, f"{c}10", f"={c}5", BLACK, False, 9, fmt=NUM, align="center", bd=True)
        else:
            p = MCOLS[i-1]
            s(ws, f"{c}10", f"=MAX(0,{c}5-{p}5)", BLACK, False, 9, fmt=NUM, align="center", bd=True)
        # 11 marketing
        s(ws, f"{c}11", f"={c}10*{cac}", BLACK, False, 10, fmt=THB, align="center", bd=True)
        # 12 operating (green link)
        s(ws, f"{c}12", f"={opc}", GREEN, False, 10, fmt=THB, align="center", bd=True)
        # 13 total cost
        s(ws, f"{c}13", f"={c}11+{c}12", BLACK, False, 10, fmt=THB, align="center", bd=True)
        # 14 net profit
        s(ws, f"{c}14", f"={c}9-{c}13", BLACK, True, 10, fmt=THB, align="center", bd=True, fill=YEL_FILL)
        # 15 cumulative
        if i==0:
            s(ws, f"{c}15", f"=-{pre}+{c}14", BLACK, True, 10, fmt=THB, align="center", bd=True, fill=YEL_FILL)
        else:
            p = MCOLS[i-1]
            s(ws, f"{c}15", f"={p}15+{c}14", BLACK, True, 10, fmt=THB, align="center", bd=True, fill=YEL_FILL)

    s(ws, "A17", f"เริ่มต้นด้วยเงินลงทุน Pre-Launch = ", "808080", italic=True, size=9)
    s(ws, "B17", f"=-{pre}", "808080", italic=True, size=9, fmt=THB)

    ws.column_dimensions["A"].width = 26
    for c in MCOLS: ws.column_dimensions[c].width = 11
    return ws

build_scenario("Conservative", "B", "C0504D")
build_scenario("Base", "C", "1F4E78")
build_scenario("Optimistic", "D", "2E7D32")

# ============================================================ SUMMARY
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
s(sm, "A1", "Summary — เปรียบเทียบ 3 Scenario (Year 1)", WHITE, True, 14, HDR_FILL)
for col in "ABCD": sm[f"{col}1"].fill = HDR_FILL

s(sm, "A3", "Metric", WHITE, True, 11, PatternFill("solid", fgColor="404040"))
for col,name,clr in [("B","Conservative","C0504D"),("C","Base","1F4E78"),("D","Optimistic","2E7D32")]:
    s(sm, f"{col}3", name, WHITE, True, 11, PatternFill("solid", fgColor=clr), align="center")

sheets = {"B":"Conservative","C":"Base","D":"Optimistic"}
scen_col_in_assump = {"B":"B","C":"C","D":"D"}

metrics = [
    ("Pre-Launch Capital (ลงทุนแรก)", lambda sh,ac: f"=Assumptions!{ac}{ROWMAP['prelaunch']}", THB),
    ("Max Cash Drawdown (เงินติดลบสูงสุด)", lambda sh,ac: f"=MIN('{sh}'!B15:M15)", THB),
    ("Peak Paying Users", lambda sh,ac: f"=MAX('{sh}'!B7:M7)", NUM),
    ("Year-1 Gross Revenue", lambda sh,ac: f"=SUM('{sh}'!B8:M8)", THB),
    ("Year-1 Net Revenue (หลัง fee)", lambda sh,ac: f"=SUM('{sh}'!B9:M9)", THB),
    ("Year-1 Operating+Mktg Cost", lambda sh,ac: f"=SUM('{sh}'!B13:M13)", THB),
    ("Year-1 Net Profit (หลัง Pre-Launch)", lambda sh,ac: f"='{sh}'!M15", THB),
    ("Net Margin %", lambda sh,ac: f"=IFERROR('{sh}'!M15/SUM('{sh}'!B8:M8),0)", PCT),
    ("Peak Monthly Net Revenue", lambda sh,ac: f"=MAX('{sh}'!B9:M9)", THB),
    ("Break-even (เดือนที่เงินสด > 0)", lambda sh,ac:
        f"=IFERROR(INDEX('{sh}'!$B$3:$M$3,MATCH(1,INDEX(--('{sh}'!$B$15:$M$15>0),0),0)),\"ไม่คืนทุนในปี1\")", None),
]
r = 4
for label, fn, fmt in metrics:
    hl = label.startswith("Year-1 Net Profit") or label.startswith("Net Margin")
    s(sm, f"A{r}", label, BLACK, hl, 10, fill=YEL_FILL if hl else None, bd=True)
    for col in "BCD":
        sh = sheets[col]; ac = scen_col_in_assump[col]
        s(sm, f"{col}{r}", fn(sh,ac), BLACK, hl, 10, fmt=fmt, align="center", bd=True,
          fill=YEL_FILL if hl else None)
    r += 1

s(sm, "A16", "หมายเหตุ:", BLACK, True, 9)
notes = [
    "• Base = สมมติฐานที่แนะนำให้ใช้วางแผนจริง (conversion 20%, organic CAC, reverse trial)",
    "• ปรับตัวเลขทั้งหมดได้ที่ชีต Assumptions (เซลล์สีน้ำเงิน) แล้วทุกชีตอัปเดตตาม",
    "• Pre-Launch Capital ลดมาก: seed bank gen ผ่าน Claude Code flat-rate (≈0) + verify มือ → เหลือแค่ infra/misc",
    "• ตัด free 2 เดือนทิ้ง → reverse trial เก็บเงินตั้งแต่เดือน 1 (Free Period = 0)",
    "• Claude Code subscription (fixed) อยู่ใน Monthly Operating Cost แล้ว ไม่ใช่ Pre-Launch",
    "• Max Cash Drawdown = เงินที่ต้องมีในกระเป๋าจริงก่อนถึงจุดคืนทุน",
    "• ตัวเลขนี้ยังไม่รวมเงินเดือน/ค่าเสียโอกาสของ solo founder",
]
for i,n in enumerate(notes):
    s(sm, f"A{17+i}", n, "808080", italic=True, size=9)

sm.column_dimensions["A"].width = 34
for col in "BCD": sm.column_dimensions[col].width = 16

# reorder: Summary, Assumptions, Conservative, Base, Optimistic
order = ["Summary","Assumptions","Conservative","Base","Optimistic"]
wb._sheets.sort(key=lambda x: order.index(x.title))

out = r"C:\Users\User\OneDrive\เดสก์ท็อป\TCAS AI Practice App\TCAS Financial Model.xlsx"
wb.save(out)
print("saved", out)
